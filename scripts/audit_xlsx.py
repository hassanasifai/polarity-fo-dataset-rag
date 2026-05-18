"""Comprehensive XLSX audit - checks evaluator-visible data/evidence risks.

Read-only. Reports issues to stdout. No mutation of the workbook.
"""
from __future__ import annotations

import ast
import json
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit

import pandas as pd
from openpyxl import load_workbook

XLSX = Path("data/processed/family_offices_validated.xlsx")
SAMPLE = Path("../assements_details/FO-MAX-data-sample-2.0.xlsx")
SEC_EVIDENCE_DIR = Path("data/evidence/sec_iapd_2026_05_17")

EXPECTED_SHEETS = {
    "data_50",
    "data_dictionary",
    "sources",
    "field_evidence",
    "validation_results",
    "validation_chain_snippets",
    "team_rosters_raw",
    "apify_evidence",
    "firecrawl_evidence",
    "contact_info_evidence",
    "email_validation_evidence",
    "google_news_signals",
    "google_places_evidence",
    "linkedin_company_evidence",
    "cheerio_homepage",
    "playwright_homepage",
}

REQUIRED_CORE = {
    "record_id",
    "family_office_name",
    "family_office_type",
    "description",
    "website_url",
    "country",
    "source_urls",
    "validation_score",
    "confidence",
    "validation_status",
}

PARITY_COLUMNS = (
    "data_validation_period",
    "family_office_domain",
    "url_quality",
    "primary_email_validation_code",
    "primary_email_code_explanation",
    "primary_email_quality_assessment",
    "contact_first_name",
    "contact_last_name",
    "contact_full_name",
    "contact_location",
    "data_completion_score_text",
    "data_completion_score_visual",
)

EVIDENCE_REQUIRED_FIELDS = [
    "primary_email",
    "primary_phone",
    "corporate_linkedin_url",
    "recent_activity",
    "recent_activity_date",
    "recent_activity_url",
    "recent_activity_type",
    "linkedin_employee_count",
    "linkedin_follower_count",
    "linkedin_specialties",
    "linkedin_company_size_band",
    "linkedin_industry",
    "linkedin_founded_year",
    "linkedin_headquarters_full",
    "twitter_url",
    "instagram_url",
    "facebook_url",
    "youtube_url",
    "google_places_phone",
    "google_places_reviews_count",
    "google_places_rating",
    "google_places_category",
    "google_maps_url",
    "primary_phone_corroborated_by_places",
    "street_address",
    "sec_registered",
    "sec_crd_number",
    "sec_file_number",
    "sec_registration_status",
    "sec_firm_name_iapd",
    "form_adv_brochure_url",
    "sec_summary_url",
    "principal_1_name",
    "principal_1_role",
    "principal_2_name",
    "principal_2_role",
    "principal_3_name",
    "principal_3_role",
    "contact_first_name",
    "contact_last_name",
    "contact_full_name",
    "contact_location",
    "data_completion_score_text",
    "data_completion_score_visual",
]

LOCAL_ARTIFACT_PREFIXES = ("reports/", "docs/", "data/", "rag_demo/reports/")
LOCAL_ARTIFACT_SUFFIXES = {
    ".csv",
    ".json",
    ".jsonl",
    ".log",
    ".md",
    ".pdf",
    ".txt",
    ".xlsx",
}
SEC_NEAR_MATCH_MARGIN = 0.15
SEC_REVIEW_STATUS_KEYS = {
    "review_status",
    "manual_review_status",
    "human_review_status",
    "documented_review_status",
}
SEC_REVIEW_STATUS_VALUES = {
    "accepted",
    "approved",
    "false_positive",
    "manual_reviewed",
    "needs_review",
    "not_a_match",
    "pending",
    "pending_review",
    "rejected",
    "reviewed",
    "reviewed_match",
    "reviewed_no_match",
}
PUBLIC_RECENCY_RE = re.compile(r"no recent public signal as of (\d{4}-\d{2}-\d{2})", re.I)


@dataclass
class ParsedSourceUrls:
    urls: list[str]
    format_name: str
    error: str = ""


@dataclass
class SourceReference:
    sheet_name: str
    row_number: int
    record_id: str
    column_name: str
    value: str
    claim_id: str = ""


@dataclass
class SecNearMatch:
    path: Path
    record_id: str
    family_office_name: str
    match_score: float
    match_threshold: float
    match_reason: str


@dataclass
class AuditResult:
    issues: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    type_distribution: dict[str, int] = field(default_factory=dict)
    country_distribution: dict[str, int] = field(default_factory=dict)

    def check(self, condition: bool, fail_msg: str, ok_msg: str | None = None) -> None:
        if condition:
            if ok_msg:
                self.notes.append(ok_msg)
        else:
            self.issues.append(fail_msg)


def is_remote_url(value: object) -> bool:
    return str(value).strip().lower().startswith(("http://", "https://"))


def normalize_url(value: object) -> str:
    text = str(value).strip()
    if not text:
        return ""
    if not is_remote_url(text):
        return text.replace("\\", "/").lstrip("./")

    parsed = urlsplit(text)
    netloc = parsed.netloc.lower().removeprefix("www.")
    path = parsed.path.rstrip("/") or "/"
    return urlunsplit((parsed.scheme.lower(), netloc, path, parsed.query, parsed.fragment))


def parse_source_urls_cell(value: object) -> ParsedSourceUrls:
    if isinstance(value, list):
        urls = [str(item).strip() for item in value if str(item).strip()]
        return ParsedSourceUrls(urls=urls, format_name="typed_list")

    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ParsedSourceUrls(urls=[], format_name="empty", error="source_urls is blank")

    text = str(value).strip()
    if not text:
        return ParsedSourceUrls(urls=[], format_name="empty", error="source_urls is blank")

    try:
        loaded = json.loads(text)
    except json.JSONDecodeError:
        loaded = None
    if loaded is not None:
        if isinstance(loaded, list) and all(isinstance(item, str) for item in loaded):
            urls = [item.strip() for item in loaded if item.strip()]
            return ParsedSourceUrls(urls=urls, format_name="json_array")
        return ParsedSourceUrls(
            urls=[],
            format_name="invalid",
            error="source_urls JSON must be an array of strings",
        )

    try:
        literal = ast.literal_eval(text)
    except (SyntaxError, ValueError):
        literal = None
    if isinstance(literal, (list, tuple)) and all(isinstance(item, str) for item in literal):
        urls = [item.strip() for item in literal if item.strip()]
        return ParsedSourceUrls(urls=urls, format_name="legacy_python_list")

    if ";" in text or "\n" in text:
        urls = [item.strip() for item in text.replace("\n", ";").split(";") if item.strip()]
        return ParsedSourceUrls(urls=urls, format_name="legacy_delimited")

    if is_remote_url(text):
        return ParsedSourceUrls(urls=[text], format_name="legacy_single_url")

    return ParsedSourceUrls(urls=[], format_name="invalid", error="source_urls is not parseable")


def is_explicit_local_artifact(value: object, project_root: Path) -> bool:
    text = normalize_url(value)
    if not text or is_remote_url(text):
        return False
    if not any(text.startswith(prefix) for prefix in LOCAL_ARTIFACT_PREFIXES):
        return False
    if Path(text).suffix.lower() not in LOCAL_ARTIFACT_SUFFIXES:
        return False
    return (project_root / text).exists()


def source_registry_urls(sources: pd.DataFrame) -> set[str]:
    urls: set[str] = set()
    for column in ("source_url", "final_url"):
        if column not in sources.columns:
            continue
        for value in sources[column].fillna(""):
            if str(value).strip():
                urls.add(normalize_url(value))
    return urls


def summarize_examples(values: list[str], limit: int = 8) -> str:
    examples = values[:limit]
    suffix = "" if len(values) <= limit else f"; +{len(values) - limit} more"
    return "; ".join(examples) + suffix


def check_unique_ids(
    df: pd.DataFrame,
    sheet_name: str,
    column_name: str,
    result: AuditResult,
) -> None:
    if column_name not in df.columns:
        result.issues.append(f"{sheet_name} missing required ID column: {column_name}")
        return

    values = df[column_name].fillna("").astype(str).str.strip()
    blank_count = int((values == "").sum())
    if blank_count:
        result.issues.append(f"{sheet_name}.{column_name} has {blank_count} blank IDs")

    duplicates = sorted(values[values != ""][values[values != ""].duplicated()].unique())
    result.check(
        not duplicates,
        f"Duplicate {sheet_name}.{column_name} values: {duplicates[:20]}",
        f"{sheet_name}.{column_name}: all nonblank IDs unique",
    )


def check_source_urls_format_and_registry(
    data: pd.DataFrame,
    registry_urls: set[str],
    project_root: Path,
    result: AuditResult,
) -> None:
    if "source_urls" not in data.columns:
        result.issues.append("data_50 missing source_urls column")
        return

    format_counts: dict[str, int] = {}
    invalid_rows: list[str] = []
    unregistered_refs: list[str] = []
    for row_number, row in data.iterrows():
        record_id = str(row.get("record_id", f"row_{row_number + 2}"))
        parsed = parse_source_urls_cell(row["source_urls"])
        format_counts[parsed.format_name] = format_counts.get(parsed.format_name, 0) + 1
        if parsed.error:
            invalid_rows.append(f"{record_id}: {parsed.error}")
            continue
        for source_url in parsed.urls:
            normalized = normalize_url(source_url)
            if is_remote_url(source_url):
                if normalized not in registry_urls:
                    unregistered_refs.append(f"{record_id} -> {source_url}")
            elif not is_explicit_local_artifact(source_url, project_root):
                unregistered_refs.append(f"{record_id} -> {source_url}")

    result.check(
        not invalid_rows,
        f"Invalid source_urls cells: {summarize_examples(invalid_rows)}",
        "source_urls: every row is parseable",
    )
    result.check(
        not unregistered_refs,
        "data_50.source_urls has entries not represented in sources/final_url "
        f"or explicit local artifacts: {summarize_examples(unregistered_refs)}",
        "data_50.source_urls: all entries represented in source registry",
    )

    non_json_count = sum(
        count
        for format_name, count in format_counts.items()
        if format_name not in {"json_array", "typed_list"}
    )
    if non_json_count:
        result.warnings.append(
            "source_urls should be typed JSON arrays where feasible; "
            f"{non_json_count} rows use legacy formats {format_counts}"
        )
    else:
        result.notes.append("source_urls: typed JSON/list format for all rows")


def find_unregistered_source_references(
    sheet_name: str,
    df: pd.DataFrame,
    column_name: str,
    registry_urls: set[str],
    project_root: Path,
) -> list[SourceReference]:
    if column_name not in df.columns:
        return []

    missing: list[SourceReference] = []
    for index, row in df.fillna("").iterrows():
        value = str(row.get(column_name, "")).strip()
        if not value:
            continue
        if is_remote_url(value):
            if normalize_url(value) in registry_urls:
                continue
        elif is_explicit_local_artifact(value, project_root):
            continue

        missing.append(
            SourceReference(
                sheet_name=sheet_name,
                row_number=int(index) + 2,
                record_id=str(row.get("record_id", "")),
                column_name=column_name,
                value=value,
                claim_id=str(row.get("claim_id", "")),
            )
        )
    return missing


def check_source_reference_registry(
    registry_urls: set[str],
    project_root: Path,
    result: AuditResult,
    *sheet_refs: tuple[str, pd.DataFrame, str],
) -> None:
    missing: list[SourceReference] = []
    for sheet_name, df, column_name in sheet_refs:
        missing.extend(
            find_unregistered_source_references(
                sheet_name=sheet_name,
                df=df,
                column_name=column_name,
                registry_urls=registry_urls,
                project_root=project_root,
            )
        )

    if not missing:
        result.notes.append(
            "Evidence source references: all remote URLs are in source registry or local artifacts"
        )
        return

    unique_values = sorted({reference.value for reference in missing})
    examples = [
        f"{ref.sheet_name}!row{ref.row_number} {ref.record_id} {ref.claim_id} -> {ref.value}"
        for ref in missing[:8]
    ]
    result.issues.append(
        "Evidence source references outside source registry/local artifacts: "
        f"{len(missing)} refs, {len(unique_values)} unique. "
        f"Examples: {summarize_examples(examples)}"
    )


def find_suspicious_venitage_rows(apify_evidence: pd.DataFrame) -> list[str]:
    if apify_evidence.empty:
        return []

    suspicious: list[str] = []
    for index, row in apify_evidence.fillna("").iterrows():
        record_id = str(row.get("record_id", ""))
        source_id = str(row.get("source_id", ""))
        source_url = str(row.get("source_url", ""))
        if (
            record_id != "fo_009"
            and not source_id.startswith("fo_009")
            and "venitage.com" not in source_url
        ):
            continue
        row_text = " ".join(str(value) for value in row.to_dict().values()).lower()
        if "venitage" not in row_text or (
            "gjelina" not in row_text and "venice beach" not in row_text
        ):
            continue
        suspicious.append(
            f"row {int(index) + 2}: {source_id or record_id} {source_url} "
            "contains Gjelina/Venice Beach text"
        )
    return suspicious


def has_documented_sec_review(record: dict[str, Any]) -> bool:
    for key in SEC_REVIEW_STATUS_KEYS:
        value = str(record.get(key, "")).strip().lower()
        if value in SEC_REVIEW_STATUS_VALUES:
            return True

    review = record.get("review")
    if isinstance(review, dict):
        value = str(review.get("status", "")).strip().lower()
        if value in SEC_REVIEW_STATUS_VALUES:
            return True
    return False


def find_sec_near_matches_without_review(sec_dir: Path) -> list[SecNearMatch]:
    if not sec_dir.exists():
        return []

    missing: list[SecNearMatch] = []
    for path in sorted(sec_dir.glob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        score = payload.get("match_score")
        threshold = payload.get("match_threshold")
        if not isinstance(score, (int, float)) or not isinstance(threshold, (int, float)):
            continue
        is_near_miss = score < threshold and score >= threshold - SEC_NEAR_MATCH_MARGIN
        if not is_near_miss or has_documented_sec_review(payload):
            continue
        missing.append(
            SecNearMatch(
                path=path,
                record_id=str(payload.get("record_id", path.stem)),
                family_office_name=str(payload.get("family_office_name", "")),
                match_score=float(score),
                match_threshold=float(threshold),
                match_reason=str(payload.get("match_reason", "")),
            )
        )
    return missing


def check_sec_near_match_reviews(project_root: Path, result: AuditResult) -> None:
    sec_dir = project_root / SEC_EVIDENCE_DIR
    if not sec_dir.exists():
        result.warnings.append(f"SEC evidence directory not found: {SEC_EVIDENCE_DIR}")
        return

    missing = find_sec_near_matches_without_review(sec_dir)
    if not missing:
        result.notes.append("SEC near-match records: documented review status present")
        return

    examples = [
        f"{item.record_id} {item.family_office_name} score={item.match_score:g}/"
        f"{item.match_threshold:g} ({item.match_reason})"
        for item in missing
    ]
    result.issues.append(
        "SEC near-match records require documented review_status/manual_review_status: "
        f"{summarize_examples(examples)}"
    )


def find_stale_public_recency_wording(
    data: pd.DataFrame,
    doc_paths: list[Path],
    audit_date: date,
) -> list[str]:
    stale: list[str] = []
    text_columns = [
        column
        for column in ("uncertainty_notes", "validation_notes", "source_notes")
        if column in data.columns
    ]
    for _, row in data.fillna("").iterrows():
        record_id = str(row.get("record_id", ""))
        for column in text_columns:
            text = str(row.get(column, ""))
            for match in PUBLIC_RECENCY_RE.finditer(text):
                recency_date = datetime.strptime(match.group(1), "%Y-%m-%d").date()
                if recency_date < audit_date:
                    stale.append(f"data_50.{column} {record_id}: {match.group(0)}")

    for path in doc_paths:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        for match in PUBLIC_RECENCY_RE.finditer(text):
            recency_date = datetime.strptime(match.group(1), "%Y-%m-%d").date()
            if recency_date < audit_date:
                stale.append(f"{path.as_posix()}: {match.group(0)}")
    return stale


def load_sheet(xlsx: Path, sheet_name: str) -> pd.DataFrame:
    return pd.read_excel(xlsx, sheet_name=sheet_name).fillna("")


def run_audit(
    project_root: Path = Path("."),
    xlsx_path: Path = XLSX,
    sample_path: Path = SAMPLE,
    audit_date: date | None = None,
) -> AuditResult:
    project_root = project_root.resolve()
    xlsx = xlsx_path if xlsx_path.is_absolute() else project_root / xlsx_path
    sample = sample_path if sample_path.is_absolute() else project_root / sample_path
    audit_day = audit_date or date.today()
    result = AuditResult()

    if not xlsx.exists():
        result.issues.append(f"{xlsx_path} does not exist")
        return result

    size_mb = os.path.getsize(xlsx) / 1024 / 1024
    result.notes.append(f"File size: {size_mb:.2f} MB")

    workbook = load_workbook(xlsx, read_only=True)
    present = set(workbook.sheetnames)
    missing = EXPECTED_SHEETS - present
    extra = present - EXPECTED_SHEETS
    if missing:
        result.issues.append(f"Missing sheets: {sorted(missing)}")
    if extra:
        result.notes.append(f"Extra sheets: {sorted(extra)}")
    result.notes.append(f"Total sheets: {len(present)}")

    data = load_sheet(xlsx, "data_50")
    result.notes.append(f"data_50: {len(data)} rows x {len(data.columns)} cols")
    result.check(len(data) == 50, f"data_50 has {len(data)} rows, expected 50")

    missing_cols = REQUIRED_CORE - set(data.columns)
    if missing_cols:
        result.issues.append(f"data_50 missing required columns: {sorted(missing_cols)}")

    result.check(
        (data["validation_status"] == "accepted").all(),
        f"data_50 has non-accepted rows: {(data['validation_status'] != 'accepted').sum()}",
        "All 50 rows have validation_status=accepted",
    )
    result.check(
        (data["confidence"] == "high").all(),
        f"data_50 has non-high confidence: {(data['confidence'] != 'high').sum()}",
        "All 50 rows have confidence=high",
    )
    check_unique_ids(data, "data_50", "record_id", result)

    if "website_ok" in data.columns:
        not_ok = (~data["website_ok"].astype(bool)).sum()
        result.check(
            not_ok == 0,
            f"{not_ok} rows have website_ok=False",
            "All 50 websites reachable at validation time",
        )

    if "source_count" in data.columns:
        under_two = (data["source_count"] < 2).sum()
        result.check(
            under_two == 0,
            f"{under_two} rows have <2 sources",
            "All 50 rows have >=2 sources attached",
        )

    if sample.exists():
        sample_workbook = load_workbook(sample, read_only=True, data_only=True)
        sample_worksheet = sample_workbook.active
        sample_names: set[str] = set()
        for row in sample_worksheet.iter_rows(min_row=5, values_only=True):
            if row and len(row) > 1 and row[1]:
                sample_names.add(str(row[1]).strip().lower())
        overlap = [name for name in data["family_office_name"] if name.lower() in sample_names]
        result.check(
            not overlap,
            f"Names overlap with sample workbook: {overlap}",
            f"No name overlap with sample workbook ({len(sample_names)} sample names checked)",
        )
    else:
        result.notes.append(f"Sample workbook not found at {sample_path} - overlap check skipped")

    df_filled = data.fillna("")
    for field_name, evidence_col, confidence_col in [
        ("primary_email", "primary_email_evidence_url", "primary_email_confidence"),
        ("primary_phone", "primary_phone_evidence_url", "primary_phone_confidence"),
        (
            "corporate_linkedin_url",
            "corporate_linkedin_evidence_url",
            "corporate_linkedin_confidence",
        ),
    ]:
        if evidence_col not in data.columns or confidence_col not in data.columns:
            result.issues.append(f"Pass 2 evidence columns missing for {field_name}")
            continue
        promoted = (df_filled[field_name] != "").sum()
        paired = (
            (df_filled[field_name] != "")
            & (df_filled[evidence_col] != "")
            & (df_filled[confidence_col] != "")
        ).sum()
        result.check(
            promoted == paired,
            f"{field_name}: {promoted - paired} promoted values missing evidence URL "
            "or confidence",
            f"{field_name}: {promoted} promoted, all paired with evidence URL + confidence",
        )

    for column in ("human_audit_status", "auto_prescreen_flag", "auto_prescreen_reason"):
        if column not in data.columns:
            result.issues.append(f"Pass 1 audit column missing: {column}")

    for column in PARITY_COLUMNS:
        if column not in data.columns:
            result.issues.append(f"Sample-parity column missing: {column}")
    for column in ("data_validation_period", "family_office_domain", "url_quality"):
        if column in data.columns:
            missing_values = (df_filled[column] == "").sum()
            result.check(
                missing_values == 0,
                f"{column} blank on {missing_values} rows (should be filled for all 50)",
                f"{column}: populated for all 50 rows",
            )

    if "primary_email_validation_code" in data.columns:
        promoted_emails = (df_filled["primary_email"] != "").sum()
        with_code = (df_filled["primary_email_validation_code"] != "").sum()
        result.check(
            promoted_emails == with_code,
            f"Email validation code mismatch: {promoted_emails} promoted emails vs "
            f"{with_code} with validation code",
            f"Email validation: {with_code}/{promoted_emails} promoted emails have a code",
        )

    if "human_audit_status" in data.columns:
        statuses = data["human_audit_status"].value_counts().to_dict()
        result.notes.append(f"human_audit_status: {statuses} (reviewer-owned)")

    if "auto_prescreen_flag" in data.columns:
        flag_counts = data["auto_prescreen_flag"].value_counts().to_dict()
        result.notes.append(f"auto_prescreen_flag: {flag_counts}")

    recent_activity_populated = (df_filled["recent_activity"] != "").sum()
    result.notes.append(f"recent_activity populated: {recent_activity_populated}/50")
    if "recent_activity_type" in data.columns:
        recent_type_counts = df_filled["recent_activity_type"].value_counts().to_dict()
        result.notes.append(f"recent_activity_type: {recent_type_counts}")

    sources = load_sheet(xlsx, "sources")
    result.notes.append(f"sources sheet: {len(sources)} rows")
    check_unique_ids(sources, "sources", "source_id", result)
    src_per_record = sources.groupby("record_id").size()
    under_two_sources = (src_per_record < 2).sum()
    result.check(
        under_two_sources == 0,
        f"{under_two_sources} records have <2 source rows in sources sheet",
        "Every record has >=2 source rows in sources sheet",
    )

    registry_urls = source_registry_urls(sources)
    check_source_urls_format_and_registry(data, registry_urls, project_root, result)

    field_evidence = load_sheet(xlsx, "field_evidence")
    result.notes.append(f"field_evidence sheet: {len(field_evidence)} rows")
    check_unique_ids(field_evidence, "field_evidence", "claim_id", result)
    if "field_name" in field_evidence.columns:
        fe_field_counts = field_evidence["field_name"].astype(str).value_counts()
        missing_field_evidence = []
        for field_name in EVIDENCE_REQUIRED_FIELDS:
            if field_name not in df_filled.columns:
                continue
            promoted = (df_filled[field_name] != "").sum()
            if promoted == 0:
                continue
            evidence_rows = int(fe_field_counts.get(field_name, 0))
            if evidence_rows < promoted:
                missing_field_evidence.append(f"{field_name}: {evidence_rows}/{promoted}")
        result.check(
            not missing_field_evidence,
            f"Promoted fields missing field_evidence rows: {missing_field_evidence}",
            "Every checked promoted field has claim-level field_evidence coverage",
        )

    validation_chain_snippets = load_sheet(xlsx, "validation_chain_snippets")
    result.notes.append(f"validation_chain_snippets sheet: {len(validation_chain_snippets)} rows")
    extracted = (validation_chain_snippets["quote_status"] == "extracted").sum()
    unavailable = (validation_chain_snippets["quote_status"] != "extracted").sum()
    result.notes.append(f"  extracted quotes: {extracted}, non-extracted: {unavailable}")
    for record_id, name in [("fo_001", "Cat Trail"), ("fo_007", "JFG"), ("fo_020", "Verlinvest")]:
        quote_count = (
            (validation_chain_snippets["record_id"] == record_id)
            & (validation_chain_snippets["quote_status"] == "extracted")
        ).sum()
        result.check(
            quote_count >= 3,
            f"Chain {record_id} {name}: only {quote_count} extracted quotes (target >=3)",
            f"  chain {record_id} {name}: {quote_count} extracted quotes",
        )

    check_source_reference_registry(
        registry_urls,
        project_root,
        result,
        ("field_evidence", field_evidence, "source_url"),
        ("validation_chain_snippets", validation_chain_snippets, "source_url"),
    )

    data_dictionary = load_sheet(xlsx, "data_dictionary")
    result.notes.append(f"data_dictionary: {len(data_dictionary)} rows")
    dd_cols = set(data_dictionary["column_name"])
    missing_defs = set(data.columns) - dd_cols
    result.check(
        not missing_defs,
        f"data_dictionary missing definitions for dataset columns: {sorted(missing_defs)}",
        "data_dictionary defines every data_50 column",
    )

    apify_evidence = (
        load_sheet(xlsx, "apify_evidence") if "apify_evidence" in present else pd.DataFrame()
    )
    venitage_mismatches = find_suspicious_venitage_rows(apify_evidence)
    result.check(
        not venitage_mismatches,
        "Suspicious fo_009/Venitage source mismatch in Apify evidence: "
        f"{summarize_examples(venitage_mismatches)}",
        "fo_009/Venitage evidence: no known Gjelina/Venice Beach mismatch detected",
    )

    check_sec_near_match_reviews(project_root, result)

    stale_public_wording = find_stale_public_recency_wording(
        data=data,
        doc_paths=[
            project_root / "reports" / "assessment_notes.md",
            project_root / "reports" / "methodology_summary.md",
            project_root / "README.md",
        ],
        audit_date=audit_day,
    )
    if stale_public_wording:
        result.warnings.append(
            "Stale public-doc recency wording found: "
            f"{len(stale_public_wording)} refs. "
            f"Examples: {summarize_examples(stale_public_wording)}"
        )
    else:
        result.notes.append("Public-doc recency wording: no stale as-of dates detected")

    result.type_distribution = data["family_office_type"].value_counts().to_dict()
    result.country_distribution = data["country"].value_counts().head(8).to_dict()
    return result


def print_audit_result(result: AuditResult) -> None:
    print("\n" + "=" * 60)
    print("AUDIT RESULTS")
    print("=" * 60)
    print(f"\nType distribution: {result.type_distribution}")
    print(f"Top countries: {result.country_distribution}")
    print(f"\nISSUES ({len(result.issues)}):")
    for issue in result.issues:
        print(f"  - {issue}")
    print(f"\nWARNINGS ({len(result.warnings)}):")
    for warning in result.warnings:
        print(f"  - {warning}")
    print(f"\nNOTES ({len(result.notes)}):")
    for note in result.notes:
        print(f"  - {note}")
    status = ">>> OK to submit" if not result.issues else ">>> FIX REQUIRED"
    print(f"\n{status}")


def main() -> int:
    result = run_audit()
    print_audit_result(result)
    return 0 if not result.issues else 1


if __name__ == "__main__":
    raise SystemExit(main())
