from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import pandas as pd
from openpyxl import load_workbook

from fo_dataset_pipeline.models import ValidatedFamilyOfficeRecord


def read_records(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(path)
    df = pd.read_csv(path).fillna("")
    return df.to_dict(orient="records")


def read_sample_workbook_names(path: Path) -> set[str]:
    if not path.exists():
        return set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    header_row = None
    name_col = None
    for row in worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True):
        for index, value in enumerate(row, start=1):
            if value == "Family Office Name":
                header_row = row
                name_col = index
                break
        if header_row is not None:
            break
    if name_col is None:
        return set()
    names: set[str] = set()
    for row in worksheet.iter_rows(min_row=5, values_only=True):
        value = row[name_col - 1] if len(row) >= name_col else None
        if value:
            names.add(str(value).strip())
    return names


def _domain(url: str) -> str:
    return urlparse(url).netloc.lower().removeprefix("www.")


def _source_type(source_url: str, website_url: str) -> str:
    source_domain = _domain(source_url)
    website_domain = _domain(website_url)
    if source_domain == website_domain:
        if source_url.lower().endswith(".pdf"):
            return "official_disclosure"
        return "official_site"
    if source_domain in {
        "adviserinfo.sec.gov",
        "reports.adviserinfo.sec.gov",
        "files.adviserinfo.sec.gov",
    }:
        return "regulator"
    if source_domain in {"find-and-update.company-information.service.gov.uk"}:
        return "registry"
    if source_domain == "businesswire.com":
        return "press_release"
    return "secondary_source"


def _source_rank(source_type: str) -> str:
    if source_type in {"official_site", "official_disclosure", "regulator", "registry"}:
        return "primary"
    if source_type == "press_release":
        return "secondary"
    return "tertiary"


def build_source_registry(records: list[ValidatedFamilyOfficeRecord]) -> list[dict[str, Any]]:
    captured_at = datetime.now(UTC).date().isoformat()
    rows: list[dict[str, Any]] = []
    for record in records:
        checks_by_url = {check.url: check for check in record.source_checks}
        for index, source_url in enumerate(record.source_urls, start=1):
            source_url_text = str(source_url)
            source_type = _source_type(source_url_text, str(record.website_url))
            check = checks_by_url.get(source_url_text)
            rows.append(
                {
                    "source_id": f"{record.record_id}_src_{index:02d}",
                    "record_id": record.record_id,
                    "family_office_name": record.family_office_name,
                    "source_url": source_url_text,
                    "source_domain": _domain(source_url_text),
                    "source_type": source_type,
                    "source_rank": _source_rank(source_type),
                    "source_owner": "first_party"
                    if _domain(source_url_text) == _domain(str(record.website_url))
                    else "third_party_or_regulator",
                    "http_ok": check.ok if check else False,
                    "http_status_code": check.status_code if check else None,
                    "final_url": check.final_url if check else "",
                    "captured_at": captured_at,
                    "retrieval_method": "live_http_check",
                    "notes": record.source_notes,
                }
            )
    return rows


def build_field_evidence(records: list[ValidatedFamilyOfficeRecord]) -> list[dict[str, Any]]:
    reviewed_at = datetime.now(UTC).date().isoformat()
    evidence_fields = [
        "family_office_name",
        "family_office_type",
        "description",
        "investment_thesis",
        "investing_sectors",
        "aum_text",
        "website_url",
        "city",
        "state_region",
        "country",
        "principal_name",
        "principal_title",
        "recent_activity",
    ]
    rows: list[dict[str, Any]] = []
    for record in records:
        source_url = str(record.source_urls[0])
        for field_name in evidence_fields:
            value = getattr(record, field_name)
            if not value:
                continue
            rows.append(
                {
                    "claim_id": f"{record.record_id}_{field_name}",
                    "record_id": record.record_id,
                    "field_name": field_name,
                    "claim_value": value.value if hasattr(value, "value") else str(value),
                    "source_url": source_url,
                    "source_type": _source_type(source_url, str(record.website_url)),
                    "claim_type": "fact"
                    if field_name
                    in {
                        "family_office_name",
                        "website_url",
                        "city",
                        "state_region",
                        "country",
                        "principal_name",
                        "principal_title",
                    }
                    else "normalized_summary",
                    "validation_status": "verified",
                    "confidence_score": record.validation_score,
                    "evidence_snippet": record.source_notes,
                    "validation_logic": record.validation_notes,
                    "reviewed_at": reviewed_at,
                }
            )
    return rows


def build_validation_results(records: list[ValidatedFamilyOfficeRecord]) -> list[dict[str, Any]]:
    return [
        {
            "record_id": record.record_id,
            "family_office_name": record.family_office_name,
            "validation_status": record.validation_status,
            "confidence": record.confidence,
            "validation_score": record.validation_score,
            "website_ok": record.website_check.ok,
            "website_status_code": record.website_check.status_code,
            "sources_attached": record.source_count,
            "sources_reachable": sum(1 for check in record.source_checks if check.ok),
            "validation_notes": record.validation_notes,
            "uncertainty_notes": record.uncertainty_notes,
        }
        for record in records
    ]


def build_data_dictionary() -> list[dict[str, str]]:
    return [
        {
            "column_name": "family_office_type",
            "definition": (
                "Explicit classification; not all rows are classic single-family offices."
            ),
        },
        {
            "column_name": "source_urls",
            "definition": "Semicolon/list of public evidence URLs checked during live validation.",
        },
        {
            "column_name": "validation_score",
            "definition": "Completeness and source-quality score from 0 to 100.",
        },
        {
            "column_name": "confidence",
            "definition": "High/medium/low confidence derived from score and validation gates.",
        },
        {
            "column_name": "uncertainty_notes",
            "definition": "Human-readable caveats preserving uncertainty instead of hiding it.",
        },
    ]


def write_outputs(
    records: list[dict[str, Any]],
    output_dir: Path,
    validated_records: list[ValidatedFamilyOfficeRecord] | None = None,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(records)
    csv_path = output_dir / "family_offices_validated.csv"
    xlsx_path = output_dir / "family_offices_validated.xlsx"
    json_path = output_dir / "family_offices_validated.json"
    df.to_csv(csv_path, index=False)
    json_path.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="data_50", index=False)
        if validated_records is not None:
            source_registry = build_source_registry(validated_records)
            field_evidence = build_field_evidence(validated_records)
            validation_results = build_validation_results(validated_records)
            pd.DataFrame(source_registry).to_excel(writer, sheet_name="sources", index=False)
            pd.DataFrame(field_evidence).to_excel(
                writer, sheet_name="field_evidence", index=False
            )
            pd.DataFrame(validation_results).to_excel(
                writer, sheet_name="validation_results", index=False
            )
            pd.DataFrame(build_data_dictionary()).to_excel(
                writer, sheet_name="data_dictionary", index=False
            )
            pd.DataFrame(source_registry).to_csv(
                output_dir / "source_registry.csv", index=False
            )
            pd.DataFrame(field_evidence).to_csv(
                output_dir / "field_evidence.csv", index=False
            )
            pd.DataFrame(validation_results).to_csv(
                output_dir / "validation_results.csv", index=False
            )
            apify_evidence_path = output_dir / "apify_crawl_evidence.csv"
            if apify_evidence_path.exists():
                pd.read_csv(apify_evidence_path).to_excel(
                    writer, sheet_name="apify_evidence", index=False
                )
            optional_evidence_sheets = {
                "firecrawl_crawl_evidence.csv": "firecrawl_evidence",
                "contact_info_evidence.csv": "contact_info_evidence",
                "email_validation_evidence.csv": "email_validation_evidence",
                "google_news_recent_signals.csv": "google_news_signals",
                "google_places_evidence.csv": "google_places_evidence",
                "linkedin_company_evidence.csv": "linkedin_company_evidence",
                "cheerio_homepage_evidence.csv": "cheerio_homepage",
                "playwright_homepage_evidence.csv": "playwright_homepage",
            }
            for csv_name, sheet_name in optional_evidence_sheets.items():
                evidence_path = output_dir / csv_name
                if evidence_path.exists():
                    pd.read_csv(evidence_path).to_excel(
                        writer, sheet_name=sheet_name, index=False
                    )
