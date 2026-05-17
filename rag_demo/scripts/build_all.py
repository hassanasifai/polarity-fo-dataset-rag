from __future__ import annotations

import hashlib
import json
import shutil
import sys
from collections import Counter
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.chunking.build_chunks import build_chunks_file, metadata_key_coverage
from src.config import (
    CHUNKS_PATH,
    GOLDEN_EVAL_PATH,
    MANIFEST_PATH,
    RAW_DATA_PATH,
    REPORTS_DIR,
    UPSTREAM_JSON_PATH,
    UPSTREAM_XLSX_PATH,
    ensure_directories,
)
from src.eval.run_eval import GOLDEN_QUESTIONS, run_eval_to_report, write_golden_eval
from src.indexing.build_bm25 import build_bm25_index
from src.indexing.build_dense_index import build_dense_index
from src.loaders.family_offices import load_family_offices


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _copy_locked_dataset() -> None:
    if not UPSTREAM_JSON_PATH.exists():
        raise FileNotFoundError(f"Expected upstream validated JSON at {UPSTREAM_JSON_PATH}")
    RAW_DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(UPSTREAM_JSON_PATH, RAW_DATA_PATH)


def _validate_xlsx_parity(records: list[dict[str, Any]]) -> dict[str, Any]:
    if not UPSTREAM_XLSX_PATH.exists():
        return {"xlsx_present": False}
    try:
        import openpyxl
    except ImportError:
        return {"xlsx_present": True, "xlsx_checked": False, "reason": "openpyxl unavailable"}

    workbook = openpyxl.load_workbook(UPSTREAM_XLSX_PATH, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    data_rows = worksheet.max_row - 1
    if data_rows != len(records):
        raise ValueError(f"XLSX row count {data_rows} does not match JSON row count {len(records)}.")
    if len(headers) != len(records[0]):
        raise ValueError(
            f"XLSX column count {len(headers)} does not match JSON column count {len(records[0])}."
        )
    return {
        "xlsx_present": True,
        "xlsx_checked": True,
        "xlsx_sheet": worksheet.title,
        "xlsx_rows": data_rows,
        "xlsx_columns": len(headers),
        "xlsx_sha256": sha256_file(UPSTREAM_XLSX_PATH),
    }


def write_manifest(records: list[dict[str, Any]], dense_info: dict[str, Any], bm25_info: dict[str, Any]) -> Path:
    chunks = []
    with CHUNKS_PATH.open("r", encoding="utf-8") as file:
        for line in file:
            if line.strip():
                chunks.append(json.loads(line))
    chunk_type_counts = Counter(chunk["metadata"]["chunk_type"] for chunk in chunks)
    manifest = {
        "dataset_path": str(RAW_DATA_PATH),
        "dataset_sha256": sha256_file(RAW_DATA_PATH),
        "upstream_json_path": str(UPSTREAM_JSON_PATH),
        "upstream_json_sha256": sha256_file(UPSTREAM_JSON_PATH),
        "row_count": len(records),
        "column_count": len(records[0]) if records else 0,
        "chunk_count": len(chunks),
        "chunk_type_counts": dict(sorted(chunk_type_counts.items())),
        "metadata_key_coverage": metadata_key_coverage(build_chunks_file(records, CHUNKS_PATH)),
        "dense_index": dense_info,
        "bm25_index": bm25_info,
        "golden_question_count": len(GOLDEN_QUESTIONS),
        "local_only": True,
        "paid_api_required": False,
        "sensitive_field_policy": "exact-copy-only; blank means not evidenced in locked dataset",
        **_validate_xlsx_parity(records),
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return MANIFEST_PATH


def main() -> None:
    ensure_directories()
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    _copy_locked_dataset()
    records = load_family_offices(RAW_DATA_PATH)
    if len(records) != 50:
        raise ValueError(f"Expected 50 validated records, found {len(records)}.")
    if len(records[0]) != 116:
        raise ValueError(f"Expected 116 columns, found {len(records[0])}.")

    chunks = build_chunks_file(records, CHUNKS_PATH)
    bm25_info = build_bm25_index(chunks)
    dense_info = build_dense_index(chunks)
    write_golden_eval(GOLDEN_EVAL_PATH)
    manifest_path = write_manifest(records, dense_info, bm25_info)
    run_eval_to_report(GOLDEN_EVAL_PATH, REPORTS_DIR / "eval_report.md")

    print(f"Built {len(chunks)} chunks from {len(records)} records.")
    print(f"Dense index: {dense_info}")
    print(f"BM25 index: {bm25_info}")
    print(f"Manifest: {manifest_path}")
    print(f"Eval report: {REPORTS_DIR / 'eval_report.md'}")


if __name__ == "__main__":
    main()
