from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from typer.testing import CliRunner

from fo_dataset_pipeline.refresh_xlsx import app


def test_rebuild_xlsx_includes_data_dictionary_and_known_sheets(tmp_path: Path) -> None:
    processed = tmp_path / "processed"
    processed.mkdir()
    pd.DataFrame(
        [{"record_id": "fo_001", "family_office_name": "Acme",
          "family_office_type": "single_family_office", "country": "United States",
          "website_url": "https://acme.com"}]
    ).to_csv(processed / "family_offices_validated.csv", index=False)
    pd.DataFrame(
        [{"source_id": "fo_001_src_01", "record_id": "fo_001",
          "source_url": "https://acme.com"}]
    ).to_csv(processed / "source_registry.csv", index=False)

    xlsx_path = tmp_path / "out.xlsx"
    runner = CliRunner()
    result = runner.invoke(app, ["--processed-dir", str(processed),
                                 "--xlsx-output", str(xlsx_path)])
    assert result.exit_code == 0, result.output
    assert xlsx_path.exists()

    wb = load_workbook(xlsx_path, read_only=True)
    sheet_names = set(wb.sheetnames)
    assert {"data_50", "data_dictionary", "sources"} <= sheet_names


def test_rebuild_xlsx_skips_missing_optional_sheets(tmp_path: Path) -> None:
    processed = tmp_path / "processed"
    processed.mkdir()
    pd.DataFrame(
        [{"record_id": "fo_001", "family_office_name": "Acme",
          "family_office_type": "single_family_office", "country": "United States",
          "website_url": "https://acme.com"}]
    ).to_csv(processed / "family_offices_validated.csv", index=False)

    xlsx_path = tmp_path / "out.xlsx"
    runner = CliRunner()
    result = runner.invoke(app, ["--processed-dir", str(processed),
                                 "--xlsx-output", str(xlsx_path)])
    assert result.exit_code == 0, result.output
    wb = load_workbook(xlsx_path, read_only=True)
    # Optional sheets should NOT be present when their CSVs are missing
    assert "google_news_signals" not in wb.sheetnames
    assert "data_50" in wb.sheetnames


def test_rebuild_xlsx_errors_when_validated_csv_missing(tmp_path: Path) -> None:
    runner = CliRunner()
    result = runner.invoke(
        app,
        [
            "--processed-dir", str(tmp_path),
            "--xlsx-output", str(tmp_path / "x.xlsx"),
        ],
    )
    assert result.exit_code != 0
    assert "not found" in result.output.lower()
